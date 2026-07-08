import io
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

def generate_pedidos_excel(pedidos_por_fecha):
    wb = Workbook()
    
    # Estilos base
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    data_font = Font(size=10) # Fuente reducida a tamaño 10 para los datos
    
    border_side = Side(border_style="thin", color="000000")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
    
    # Colores intercalados por pedido
    fill_azul_medio = PatternFill(start_color="8DB4E2", end_color="8DB4E2", fill_type="solid")
    fill_azul_claro = PatternFill(start_color="C5D9F1", end_color="C5D9F1", fill_type="solid")
    
    if not pedidos_por_fecha:
        ws = wb.active
        ws.title = "Sin pedidos"
        ws["A1"] = "No hay pedidos para exportar"
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
    
    for fecha, pedidos in pedidos_por_fecha.items():
        if isinstance(fecha, (datetime.date, datetime.datetime)):
            sheet_name = f"Para_el_{fecha.strftime('%d-%m-%Y')}"
            title_date = fecha.strftime('%d/%m/%Y')
        else:
            sheet_name = "Sin_Fecha"
            title_date = "FECHA INDEFINIDA"
            
        ws = wb.create_sheet(title=sheet_name)
        
        # Fila 1: Título (Centrado en 12 columnas: A a L)
        ws.merge_cells("A1:L1")
        cell_title = ws["A1"]
        cell_title.value = f"PEDIDOS JOSE CARO {title_date}"
        cell_title.font = header_font
        cell_title.fill = header_fill
        cell_title.alignment = Alignment(horizontal="center", vertical="center")
        
        # Fila 2: Cabeceras
        headers = ["Nombre de negocio", "Tipo", "Direccion", "Barrio_poblacion", 
                   "Nombre cliente", "NIT", "Telefono", "Productos", 
                   "Cantidad", "Precio unit.", "Valor total", "Observaciones"]
        
        for col_num, header_title in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_num, value=header_title)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = thin_border
            
        current_row = 3
        for p_idx, pedido in enumerate(pedidos):
            cliente = pedido["cliente"]
            items = pedido["items"]
            current_fill = fill_azul_medio if p_idx % 2 == 0 else fill_azul_claro
            
            num_items = len(items)
            start_row = current_row
            end_row = current_row + num_items - 1
            
            # 1. Aplicar fondo y bordes a todo el bloque (12 columnas)
            for r in range(start_row, end_row + 1):
                for c in range(1, 13):
                    cell = ws.cell(row=r, column=c)
                    cell.fill = current_fill
                    cell.border = thin_border
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.font = data_font

            # 2. Insertar datos del cliente (Columnas 1-7)
            ws.cell(row=start_row, column=1, value=cliente.get("nombre_negocio", ""))
            ws.cell(row=start_row, column=2, value=cliente.get("tipo_negocio", ""))
            
            c_dir = ws.cell(row=start_row, column=3, value=cliente.get("direccion", ""))
            c_dir.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            ws.cell(row=start_row, column=4, value=cliente.get("barrio_poblacion", ""))
            ws.cell(row=start_row, column=5, value=cliente.get("nombre_cliente", ""))
            ws.cell(row=start_row, column=6, value=cliente.get("cc_o_nit", ""))
            ws.cell(row=start_row, column=7, value=cliente.get("telefono", ""))

            if num_items > 1:
                for col_idx in range(1, 8):
                    ws.merge_cells(start_row=start_row, start_column=col_idx, end_row=end_row, end_column=col_idx)

            # 3. Insertar productos (Columnas 8-12)
            for i, item in enumerate(items):
                r = start_row + i
                ws.cell(row=r, column=8, value=item["producto_nombre"]).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                ws.cell(row=r, column=9, value=item["cantidad"])
                
                c_precio = ws.cell(row=r, column=10, value=item["precio_unitario"])
                c_precio.number_format = '#,##0'
                
                c_total = ws.cell(row=r, column=11, value=item["cantidad"] * item["precio_unitario"])
                c_total.number_format = '#,##0'
                
                c_obs = ws.cell(row=r, column=12, value=item.get("observaciones", ""))
                c_obs.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            
            current_row = end_row + 1

        # 4. Configurar anchos fijos (Ajustados para optimizar espacio en tamaño carta)
        anchos_columnas = {
            'A': 15.0,        # Nombre de negocio
            'B': 8.0,         # Tipo
            'C': 13.0,        # Direccion
            'D': 10.0,        # Barrio_poblacion
            'E': 14.0,        # Nombre cliente
            'F': 11.0,        # NIT
            'G': 10.0,        # Telefono
            'H': 14.0,        # Productos
            'I': 7.0,         # Cantidad
            'J': 9.0,         # Precio unit.
            'K': 10.0,        # Valor total
            'L': 14.0         # Observaciones
        }
        
        for col_letter, width in anchos_columnas.items():
            ws.column_dimensions[col_letter].width = width

        # 5. Configurar impresión (Tamaño carta, ajustar a 1 página de ancho, horizontal)
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0 # Que use las páginas hacia abajo que necesite

    if len(wb.sheetnames) > 1 and "Sheet" in wb.sheetnames:
        del wb["Sheet"]
        
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_informe_excel(datos, nombre_mes, anio):
    """Genera el Excel del informe de ventas por cliente para un mes dado."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = f"{nombre_mes}_{anio}"

    # Estilos
    header_fill = PatternFill(start_color="002060", end_color="002060", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    data_font = Font(size=10)
    fill_azul_medio = PatternFill(start_color="8DB4E2", end_color="8DB4E2", fill_type="solid")
    fill_azul_claro = PatternFill(start_color="C5D9F1", end_color="C5D9F1", fill_type="solid")
    border_side = Side(border_style="thin", color="000000")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # Titulo
    ws.merge_cells("A1:D1")
    cell_title = ws["A1"]
    cell_title.value = f"INFORME DE VENTAS JOSE CARO - {nombre_mes.upper()} {anio}"
    cell_title.font = Font(color="FFFFFF", bold=True, size=12)
    cell_title.fill = header_fill
    cell_title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # Cabeceras
    headers = ["Nombre Negocio", "Nombre Cliente", "Total Ventas", "Cant. Pedidos"]
    for col_num, title in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_num, value=title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border

    # Datos
    total_ventas_general = 0
    total_pedidos_general = 0
    for idx, row in enumerate(datos):
        r = idx + 3
        current_fill = fill_azul_medio if idx % 2 == 0 else fill_azul_claro
        negocio = row.get("nombre_negocio") or row.get("cliente_nombre", "")
        cliente = row.get("cliente_nombre", "")
        ventas = row.get("total_ventas", 0)
        pedidos = row.get("cantidad_pedidos", 0)
        total_ventas_general += ventas
        total_pedidos_general += pedidos

        for c in range(1, 5):
            cell = ws.cell(row=r, column=c)
            cell.fill = current_fill
            cell.border = thin_border
            cell.font = data_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.cell(row=r, column=1, value=negocio)
        ws.cell(row=r, column=2, value=cliente)
        c_ventas = ws.cell(row=r, column=3, value=ventas)
        c_ventas.number_format = "$#,##0"
        ws.cell(row=r, column=4, value=pedidos)

    # Fila de totales
    fila_total = len(datos) + 3
    ws.merge_cells(start_row=fila_total, start_column=1, end_row=fila_total, end_column=2)
    c_lbl = ws.cell(row=fila_total, column=1, value="TOTAL GENERAL")
    c_lbl.font = Font(bold=True, size=10, color="FFFFFF")
    c_lbl.fill = header_fill
    c_lbl.alignment = Alignment(horizontal="center", vertical="center")
    c_lbl.border = thin_border

    c_tot = ws.cell(row=fila_total, column=3, value=total_ventas_general)
    c_tot.number_format = "$#,##0"
    c_tot.font = Font(bold=True, size=10, color="FFFFFF")
    c_tot.fill = header_fill
    c_tot.alignment = Alignment(horizontal="center", vertical="center")
    c_tot.border = thin_border

    c_ped = ws.cell(row=fila_total, column=4, value=total_pedidos_general)
    c_ped.font = Font(bold=True, size=10, color="FFFFFF")
    c_ped.fill = header_fill
    c_ped.alignment = Alignment(horizontal="center", vertical="center")
    c_ped.border = thin_border

    # Anchos
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14

    # Configuración de impresión
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
